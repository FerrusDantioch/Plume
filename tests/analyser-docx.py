# Analyse un fichier Word produit par les tests (appelé par export.spec.js).
# 1. Vérifie l'archive et que chaque fichier XML est bien formé.
# 2. Relit le document avec python-docx (styles, mise en forme, listes).
# 3. Le fait ouvrir par LibreOffice et convertir en PDF : preuve qu'un vrai
#    traitement de texte l'accepte. Le PDF est ensuite analysé.
# Affiche un résumé en JSON.

import json
import os
import subprocess
import sys
import tempfile
import zipfile
from xml.dom import minidom

import docx
import pdfplumber

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def format_du_morceau(run):
    f = run.font
    return {
        "texte": run.text,
        "gras": run.bold,
        "italique": run.italic,
        "souligne": bool(run.underline),
        "barre": bool(f.strike),
        "taille": f.size.pt if f.size else None,
        "police": f.name,
        "couleur": str(f.color.rgb) if f.color and f.color.type is not None else None,
        "fond": (run._r.rPr.find(W + "shd").get(W + "fill") if run._r.rPr is not None and run._r.rPr.find(W + "shd") is not None else None),
    }


def main(chemin):
    resultat = {}
    with zipfile.ZipFile(chemin) as archive:
        resultat["fichiers"] = sorted(archive.namelist())
        resultat["archive_intacte"] = archive.testzip() is None
        for nom in archive.namelist():
            if nom.endswith(".xml") or nom.endswith(".rels"):
                minidom.parseString(archive.read(nom))  # lève une erreur si mal formé

    document = docx.Document(chemin)
    resultat["titre"] = document.core_properties.title
    section = document.sections[0]
    resultat["page_cm"] = [round(section.page_width.cm, 1), round(section.page_height.cm, 1), round(section.left_margin.cm, 1)]
    paragraphes = []
    for p in document.paragraphs:
        numero = p._p.pPr.find(W + "numPr") if p._p.pPr is not None else None
        paragraphes.append({
            "style": p.style.name,
            "texte": p.text,
            "alignement": str(p.alignment) if p.alignment is not None else None,
            "liste": numero.find(W + "numId").get(W + "val") if numero is not None else None,
            "morceaux": [format_du_morceau(r) for r in p.runs if r.text],
        })
    resultat["paragraphes"] = paragraphes

    # Ouverture par LibreOffice et conversion en PDF.
    with tempfile.TemporaryDirectory() as dossier:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", dossier, chemin],
            check=True, capture_output=True, timeout=180,
            env={**os.environ, "HOME": dossier},
        )
        pdf = os.path.join(dossier, os.path.splitext(os.path.basename(chemin))[0] + ".pdf")
        with pdfplumber.open(pdf) as lecteur:
            resultat["libreoffice_texte"] = "\n".join((page.extract_text() or "") for page in lecteur.pages)
            resultat["libreoffice_pages"] = len(lecteur.pages)
        apercu = os.environ.get("APERCU_DOCX")
        if apercu:
            import pypdfium2
            rendu = pypdfium2.PdfDocument(pdf)
            rendu[0].render(scale=1.3).to_pil().save(apercu)
            rendu.close()
    print(json.dumps(resultat, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1])

# Crée les fichiers d'essai utilisés par tests/ouvrir-fichier.spec.js.
# (Python 3, python-docx et LibreOffice.) Les fichiers produits sont gardés
# dans ce dossier : les tests n'ont pas besoin de relancer ce script.

import os
import subprocess
import tempfile

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

ICI = os.path.dirname(os.path.abspath(__file__))

# 1. Un document « à la Word » (modèle par défaut de python-docx, créé par Word).
d = docx.Document()
d.core_properties.title = "Titre interne ignoré"
d.add_heading("Compte rendu", level=1)
p = d.add_paragraph("Texte ")
p.add_run("gras").bold = True
p.add_run(", ")
p.add_run("italique").italic = True
p.add_run(", ")
p.add_run("souligné").underline = True
p.add_run(", ")
r = p.add_run("barré")
r.font.strike = True
p.add_run(", ")
r = p.add_run("rouge")
r.font.color.rgb = RGBColor(0xE5, 0x48, 0x4D)
p.add_run(", ")
r = p.add_run("surligné")
r.font.highlight_color = WD_COLOR_INDEX.YELLOW
p.add_run(", ")
r = p.add_run("grand")
r.font.size = Pt(18)
p.add_run(" et ")
r = p.add_run("machine")
r.font.name = "Courier New"
p.add_run(".")
c = d.add_paragraph("Paragraphe centré")
c.alignment = WD_ALIGN_PARAGRAPH.CENTER
j = d.add_paragraph("Paragraphe justifié en double interligne")
j.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
j.paragraph_format.line_spacing = 2
d.add_heading("Listes", level=2)
d.add_paragraph("pomme", style="List Bullet")
d.add_paragraph("poire", style="List Bullet")
d.add_paragraph("un", style="List Number")
d.add_paragraph("deux", style="List Number")
l = d.add_paragraph("Ligne un")
l.add_run().add_break()
l.add_run("Ligne deux")
d.add_heading("Détails", level=3)
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text = "Case A"
t.cell(0, 1).text = "Case B"
t.cell(1, 0).text = "Case C"
t.cell(1, 1).text = "Case D"
# Une petite image (carré violet de 8 × 8 pixels), qui ne sera pas reprise.
image = os.path.join(tempfile.gettempdir(), "omni-texte-essai-image.png")
import zlib, struct
def bloc(type_, donnees):
    return struct.pack(">I", len(donnees)) + type_ + donnees + struct.pack(">I", zlib.crc32(type_ + donnees) & 0xFFFFFFFF)
lignes = b"".join(b"\x00" + b"\xa7\x77\xff" * 8 for _ in range(8))
with open(image, "wb") as f:
    f.write(b"\x89PNG\r\n\x1a\n" + bloc(b"IHDR", struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0))
            + bloc(b"IDAT", zlib.compress(lignes)) + bloc(b"IEND", b""))
d.add_picture(image)
d.add_paragraph("Fin du document.")
d.save(os.path.join(ICI, "exemple-word.docx"))

# 2. Le même document, réenregistré par LibreOffice (autre façon d'écrire le XML).
with tempfile.TemporaryDirectory() as dossier:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "docx:MS Word 2007 XML", "--outdir", dossier,
         os.path.join(ICI, "exemple-word.docx")],
        check=True, capture_output=True, timeout=180, env={**os.environ, "HOME": dossier},
    )
    with open(os.path.join(dossier, "exemple-word.docx"), "rb") as source, \
         open(os.path.join(ICI, "exemple-libreoffice.docx"), "wb") as cible:
        cible.write(source.read())

# 3. Un texte enregistré avec l'ancien encodage de Windows (accents sur un octet).
with open(os.path.join(ICI, "ancien-windows.txt"), "wb") as f:
    f.write("Première ligne : café, élève, garçon.\r\n\r\nTroisième ligne.\r\n".encode("cp1252"))

# 4. Un faux fichier Word (n'importe quoi renommé en .docx).
with open(os.path.join(ICI, "abime.docx"), "wb") as f:
    f.write(b"Ceci n'est pas un fichier Word.")
